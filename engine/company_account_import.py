"""Bounded, local-only XLSX account preview and atomic append-only import."""

from __future__ import annotations

import base64
from io import BytesIO
import re
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import load_workbook

from .receipt_grouping_api import _account_input
from .receipt_grouping_models import GroupingError, normalize_account, normalize_bank, normalize_text
from .receipt_grouping_store import GroupingStore

HEADERS = ("公司全名", "来源银行", "开户行", "完整本方账号")
KEYS = ("company_name", "bank_name", "branch_name", "account_number")
MAX_FILE_BYTES = 512 * 1024
MAX_ROWS = 500
MAX_SOURCE_CHARACTERS = 256
SOURCE_TRUNCATION_NOTICE = "…（内容过长，请回原表核对）"


class AccountImportInvalid(GroupingError):
    code = "account_import_invalid"


class AccountImportConflict(GroupingError):
    code = "account_import_conflict"


def _key(account):
    return normalize_bank(account["bank_name"]), normalize_account(account["account_number"])


def _same(left, right):
    return all(normalize_text(left.get(key) or "") == normalize_text(right.get(key) or "") for key in KEYS)


def _existing(store):
    result = {}
    offset = 0
    while True:
        page = store.account_list(active_only=False, offset=offset, limit=50)
        for account in page["items"]:
            result.setdefault(_key(account), []).append(account)
        if page["next_offset"] is None:
            return result
        offset = page["next_offset"]


def _check_duplicate(account, known):
    matches = known.get(_key(account), [])
    if not matches:
        return "ready", "可导入"
    if all(_same(account, match) for match in matches):
        return "duplicate", "已有相同档案或本表重复（含停用档案），将跳过"
    return "error", "同银行、同账号已有不同资料，请先核对或编辑原档案"


def _source_text(value):
    """Bounded display-only cell text; never use this value to save an account."""
    text = "" if value is None else str(value)
    text = text.replace("\x00", "\ufffd")
    if len(text) > MAX_SOURCE_CHARACTERS:
        return text[:MAX_SOURCE_CHARACTERS] + SOURCE_TRUNCATION_NOTICE
    return text


def _workbook(encoded):
    if not isinstance(encoded, str) or len(encoded) > ((MAX_FILE_BYTES + 2) // 3) * 4:
        raise AccountImportInvalid("请选择不超过512KB的XLSX模板")
    try:
        content = base64.b64decode(encoded, validate=True)
        if not content or len(content) > MAX_FILE_BYTES:
            raise ValueError()
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if (len(entries) > 100 or sum(entry.file_size for entry in entries) > 8 * 1024 * 1024
                    or any(entry.flag_bits & 1 for entry in entries)
                    or any("vbaproject" in entry.filename.lower() or "externallinks/" in entry.filename.lower()
                           for entry in entries)):
                raise ValueError()
            nodes = 0
            for entry in entries:
                if not entry.filename.lower().endswith((".xml", ".rels")):
                    continue
                xml = archive.read(entry)
                # XLSX needs no DTD/entities. Reject them before either XML
                # parser sees input, including hidden sheets and styles.
                # Only UTF-8 XML is accepted by this bounded template importer;
                # UTF-16/32 NUL bytes must not bypass the declaration check.
                xml_text = xml.decode("utf-8-sig")
                if "\x00" in xml_text or re.search(r"<!\s*(?:DOCTYPE|ENTITY)\b", xml_text, re.I):
                    raise ValueError()
                worksheet = entry.filename.startswith("xl/worksheets/")
                for _, element in ElementTree.iterparse(BytesIO(xml), events=("start",)):
                    nodes += 1
                    if nodes > 100000:
                        raise ValueError()
                    if worksheet and element.tag.endswith("}c"):
                        address = element.get("r", "")
                        if not re.fullmatch(r"[A-D][1-9][0-9]{0,2}", address) or int(address[1:]) > MAX_ROWS + 1:
                            raise ValueError()
        return load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise AccountImportInvalid("模板文件无效，请重新下载XLSX模板填写") from exc


def preview_accounts(store, encoded):
    workbook = _workbook(encoded)
    try:
        sheet = workbook["本方账户"] if "本方账户" in workbook.sheetnames else workbook.worksheets[0]
        if (sheet.max_column or 0) > 4 or (sheet.max_row or 0) > MAX_ROWS + 1:
            raise AccountImportInvalid("每次最多500行且仅使用模板四列，请清除多余行列")
        # Ignore claimed worksheet dimensions; bound actual iteration below.
        sheet.reset_dimensions()
        rows = sheet.iter_rows(max_col=5)
        header = next(rows, ())
        if tuple(cell.value for cell in header[:4]) != HEADERS or header[4].value is not None:
            raise AccountImportInvalid("模板列名不匹配，请使用下载的四列模板")
        known = _existing(store)
        result = []
        for number, cells in enumerate(rows, 2):
            if number > MAX_ROWS + 1:
                raise AccountImportInvalid("每次最多导入500行档案")
            if not any(cell.value is not None and cell.value != "" for cell in cells):
                continue
            source = {key: _source_text(cell.value) for key, cell in zip(KEYS, cells[:4])}
            account = None
            try:
                if cells[4].value is not None or any(cell.data_type == "f" for cell in cells):
                    raise ValueError("请使用四列文本，不要使用公式")
                if not isinstance(cells[3].value, str):
                    raise ValueError("完整本方账号必须为文本，请按模板填写以保留前导零和全部位数")
                values = [cell.value if cell.value is not None else "" for cell in cells[:4]]
                try:
                    account = _account_input(dict(zip(KEYS, values)))
                except GroupingError:
                    raise ValueError("公司全名、来源银行和完整账号必填，账号仅可含数字") from None
                status, message = _check_duplicate(account, known)
                if status == "ready":
                    known.setdefault(_key(account), []).append(account)
            except ValueError as exc:
                status, message = "error", str(exc)
            result.append({"row_number": number, "account": account, "source": source,
                           "status": status, "message": message})
        if not result:
            raise AccountImportInvalid("模板尚未填写账户资料")
        return {"rows": result, "counts": {kind: sum(row["status"] == kind for row in result)
                                          for kind in ("ready", "duplicate", "error")}}
    except AccountImportInvalid:
        raise
    except Exception as exc:
        raise AccountImportInvalid("模板内容无效，请重新下载模板填写") from exc
    finally:
        workbook.close()


def import_accounts(store, accounts):
    if not isinstance(accounts, list) or not 1 <= len(accounts) <= MAX_ROWS:
        raise AccountImportInvalid("每次仅可导入1至500行有效档案")
    try:
        validated = [_account_input(account) for account in accounts]
    except GroupingError as exc:
        raise AccountImportInvalid("档案内容无效，请重新检查模板") from exc
    created = skipped = 0
    with store.transaction():
        known = _existing(store)
        for account in validated:
            status, _ = _check_duplicate(account, known)
            if status == "error":
                raise AccountImportConflict("档案已变化或存在冲突，请重新预览模板")
            if status == "duplicate":
                skipped += 1
                continue
            saved = store.account_save(account)
            known.setdefault(_key(saved), []).append(saved)
            created += 1
    return {"created": created, "skipped": skipped}


def dispatch_account_import(op, data, database):
    with GroupingStore(database) as store:
        if op == "batch_counterparty_account_import_preview":
            return preview_accounts(store, data["workbook_base64"])
        return import_accounts(store, data["accounts"])
