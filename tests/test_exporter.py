from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

from engine.exporter import DEFAULT_HEADERS, export_bundle_index, export_index
from engine.xlsx_writer import write_xlsx_sheets


def test_export_headers_preserve_legacy_prefix_and_append_audit_fields() -> None:
    assert DEFAULT_HEADERS[:11] == (
        "source_file", "source_sha256", "source_page", "segment_no", "query",
        "matched_field", "crop_mode", "review_status", "confidence", "output_file", "processed_at",
    )
    assert DEFAULT_HEADERS[11:] == (
        "matched_keywords", "matched_text", "crop_x0", "crop_y0", "crop_x1", "crop_y1",
    )


def test_export_index_writes_readable_xlsx_with_stable_headers(tmp_path: Path) -> None:
    output = export_index(tmp_path / "nested" / "match-index.xlsx", [{
        "source_file": "sample.pdf",
        "source_page": 3,
        "query": "手续费",
        "confidence": 0.98,
    }])
    assert output.exists()
    with ZipFile(output) as archive:
        names = set(archive.namelist())
        assert "xl/workbook.xml" in names
        assert "xl/worksheets/sheet1.xml" in names
        sheet = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
        assert "手续费" in sheet
        assert "source_page" in sheet


def test_export_index_contains_multi_search_evidence(tmp_path: Path) -> None:
    output = export_index(tmp_path / "multi.xlsx", [{
        "query": "包含全部：示例实业、华夏银行；排除任一：退款",
        "matched_keywords": "示例实业、华夏银行",
        "matched_text": "交易对手：示例实业；开户行：华夏银行",
        "crop_x0": 0,
        "crop_y0": 20,
        "crop_x1": 595,
        "crop_y1": 300,
    }])

    with ZipFile(output) as archive:
        sheet = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")

    assert "matched_keywords" in sheet
    assert "示例实业、华夏银行" in sheet
    assert "交易对手：示例实业；开户行：华夏银行" in sheet
    assert "crop_y1" in sheet


def test_write_xlsx_sheets_preserves_inline_formula_text_and_sheet_order(tmp_path: Path) -> None:
    output = write_xlsx_sheets(
        tmp_path / "multi.xlsx",
        (
            ("索引", ("value",), [{"value": "=SUM(A1:A2)"}]),
            ("输出映射", ("value",), [{"value": "literal"}]),
        ),
    )

    import openpyxl

    workbook = openpyxl.load_workbook(output, data_only=False)
    assert workbook.sheetnames == ["索引", "输出映射"]
    assert workbook["索引"]["A2"].value == "=SUM(A1:A2)"
    assert workbook["索引"]["A2"].data_type == "s"
    assert workbook["输出映射"]["A2"].value == "literal"


def test_export_bundle_index_writes_three_named_tables_with_scope_item_value(tmp_path: Path) -> None:
    output = export_bundle_index(
        tmp_path / "bundle.xlsx",
        [{"source_file": "source.pdf", "source_page": 2}],
        [{"segment_id": "seg-1", "output_page": 1}],
        [{"item": "output_mode", "value": "both"}],
    )

    import openpyxl

    workbook = openpyxl.load_workbook(output, data_only=False)
    assert workbook.sheetnames == ["索引", "输出映射", "导出范围"]
    assert [cell.value for cell in workbook["导出范围"][1]] == ["item", "value"]
    assert [cell.value for cell in workbook["导出范围"][2]] == ["output_mode", "both"]


def test_export_bundle_index_round_trips_legacy_columns_and_page_mapping(tmp_path: Path) -> None:
    row = {header: None for header in DEFAULT_HEADERS}
    row.update(source_file="source.pdf", source_page=4, matched_text="=SUM(A1:A2)")
    output = export_bundle_index(
        tmp_path / "bundle-full.xlsx",
        [row],
        [{
            "segment_id": "seg-1",
            "source_file": "source.pdf",
            "source_page": 4,
            "segment_no": 2,
            "output_file": "全部匹配结果.pdf",
            "output_page": 7,
        }],
        [{"item": "selected_count", "value": 1}],
    )

    import openpyxl

    workbook = openpyxl.load_workbook(output, data_only=False)
    index_sheet = workbook["索引"]
    assert index_sheet.max_column == 17
    assert [cell.value for cell in index_sheet[1]] == list(DEFAULT_HEADERS)
    assert index_sheet["M2"].value == "=SUM(A1:A2)"
    assert index_sheet["M2"].data_type == "s"
    assert [cell.value for cell in workbook["输出映射"][2]] == [
        "seg-1", "source.pdf", 4, 2, "全部匹配结果.pdf", 7,
    ]


def test_bundle_index_removes_all_xml_10_forbidden_codepoints(tmp_path: Path) -> None:
    from io import BytesIO
    import openpyxl

    output = export_bundle_index(tmp_path / "xml-safe.xlsx", [{
        "source_file": "safe.pdf", "matched_text": "text\x00\ufffe\uffff\ud800\udfff\t\n\r\U0001f4c4",
    }], [], [])
    workbook = openpyxl.load_workbook(BytesIO(output.read_bytes()))
    try:
        # XML parsers normalize literal carriage returns to line feeds.
        assert workbook["索引"]["M2"].value == "text_____\t\n\n\U0001f4c4"
    finally:
        workbook.close()


def test_bundle_index_writes_to_caller_owned_stream_without_reopening_or_closing(tmp_path: Path) -> None:
    from io import BytesIO
    import openpyxl

    target = tmp_path / "exclusive.xlsx"
    with target.open("xb+") as stream:
        assert export_bundle_index(stream, [{"source_file": "source.pdf"}], [], []) is stream
        assert not stream.closed
        stream.flush()
        stream.seek(0)
        workbook = openpyxl.load_workbook(BytesIO(stream.read()))
        assert workbook.sheetnames == ["索引", "输出映射", "导出范围"]
        assert workbook["索引"]["A2"].value == "source.pdf"
        workbook.close()
