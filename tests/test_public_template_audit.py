from __future__ import annotations

from hashlib import sha256
import importlib.util
from pathlib import Path
import sys

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "audit-public-tree.py"
TEMPLATE_RELATIVE = "public/templates/本方账户导入模板.xlsx"
TEMPLATE = ROOT / TEMPLATE_RELATIVE


def load_auditor():
    spec = importlib.util.spec_from_file_location("audit_public_tree", SCRIPT)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def test_reviewed_template_is_the_exact_empty_four_column_asset_and_passes() -> None:
    auditor = load_auditor()
    expected_sha256 = auditor.PUBLIC_BINARY_SHA256_ALLOWLIST[TEMPLATE_RELATIVE]
    assert sha256(TEMPLATE.read_bytes()).hexdigest() == expected_sha256
    assert auditor.audit_file(TEMPLATE_RELATIVE, TEMPLATE, []) == []

    workbook = load_workbook(TEMPLATE, read_only=True, data_only=False)
    try:
        assert workbook.sheetnames == ["本方账户", "填写说明"]
        account_sheet = workbook["本方账户"]
        assert account_sheet.max_row == 501
        assert account_sheet.max_column == 4
        assert tuple(account_sheet.iter_rows(min_row=1, max_row=1, values_only=True))[0] == (
            "公司全名", "来源银行", "开户行", "完整本方账号"
        )
        assert all(
            row == (None, None, None, None)
            for row in account_sheet.iter_rows(min_row=2, max_row=501, max_col=4, values_only=True)
        )
        assert workbook["填写说明"].max_column == 1
    finally:
        workbook.close()


def test_tampering_with_the_allowlisted_path_is_rejected(tmp_path: Path) -> None:
    auditor = load_auditor()
    tampered = tmp_path / TEMPLATE.name
    data = bytearray(TEMPLATE.read_bytes())
    data[-1] ^= 1
    tampered.write_bytes(data)

    findings = auditor.audit_file(TEMPLATE_RELATIVE, tampered, [])

    assert {finding["rule"] for finding in findings} == {"private_file_type", "unexpected_binary"}


def test_other_xlsx_remains_rejected_even_with_the_same_bytes(tmp_path: Path) -> None:
    auditor = load_auditor()
    other = tmp_path / "other.xlsx"
    other.write_bytes(TEMPLATE.read_bytes())

    findings = auditor.audit_file("public/templates/other.xlsx", other, [])

    assert {finding["rule"] for finding in findings} == {"private_file_type", "unexpected_binary"}
