"""Frozen multi-PDF previews use synthetic tasks and host-owned paths."""

from copy import deepcopy
from hashlib import sha256
from pathlib import Path
from uuid import uuid4

import pymupdf
import pytest

from engine import engine as core
from engine import export_journal as journal_module
from engine import export_publish
from engine.batch_previews import register_preview
from engine.batch_store import BatchStore
from engine.export_bundle import ExportBundleService
from engine.export_scope import ExportScopeError
from test_export_scope import duplicate_export_task, export_task, save  # shared isolated fixture


@pytest.fixture
def bundle(export_task, monkeypatch):
    batch, review, _, _, records, _ = export_task
    save(export_task, records)
    root = batch.parent / "export-intents"
    previews = batch.parent / "export-previews"
    root.mkdir()
    previews.mkdir()
    monkeypatch.setattr(core, "EXPORT_OWNERSHIP_DIR", batch.parent / ".export-ownership")
    return ExportBundleService(batch, review, root, previews)


@pytest.fixture
def duplicate_bundle(duplicate_export_task, monkeypatch):
    batch, review, _, _, records, _ = duplicate_export_task
    save(duplicate_export_task, records)
    root = batch.parent / "export-intents"
    previews = batch.parent / "export-previews"
    root.mkdir()
    previews.mkdir()
    monkeypatch.setattr(core, "EXPORT_OWNERSHIP_DIR", batch.parent / ".export-ownership")
    return ExportBundleService(batch, review, root, previews)


def request_all(task, mode="both", xlsx=True):
    request = deepcopy(task[5])
    request.update(scope_kind="all", output_mode=mode, include_xlsx=xlsx,
                   selected_segment_ids=[r["id"] for r in task[4]],
                   expected_records=[{"id": r["id"], "record_revision": 1} for r in task[4]])
    return request


def register(service, created):
    with BatchStore(service.batch_database) as store:
        for file in created["files"]:
            register_preview(store, created["job_id"], file["preview_token"], service.preview_root)


def _compressed_receipt(bundle, created_at, *, blocked=False):
    intent_id = str(uuid4())
    token = str(uuid4())
    preview_path = bundle.preview_root / f"{token}.pdf"
    entry = {
        "schema": 1,
        "intent_id": intent_id,
        "job_id": "receipt-retention-job",
        "created_at": created_at,
        "state": "published",
        "files": [{"file_id": f"receipt-file-{intent_id}", "preview_token": token,
                   "preview_path": str(preview_path)}],
        "receipt": {"intent_id": intent_id},
    }
    if blocked:
        entry["residual_attempts"] = [{"path": str(preview_path)}]
        preview_path.write_bytes(b"foreign completed output")
    bundle.journal.create(entry)
    return entry


@pytest.mark.parametrize("mode,count,total", [("merged", 1, 2), ("by_source", 2, 2), ("both", 3, 4)])
def test_real_frozen_preview_modes(bundle, export_task, mode, count, total):
    before = [sha256(path.read_bytes()).hexdigest() for path in export_task[3]]
    created = bundle.create(request_all(export_task, mode))
    assert created["state"] == "created"
    assert len(created["files"]) == count
    register(bundle, created)
    preview = bundle.render(created["intent_id"])
    assert preview["state"] == "rendered"
    assert preview["total_pages"] == total
    assert len(preview["files"]) == count
    assert "records" not in preview and "evidence_by_id" not in preview
    for file in preview["files"]:
        with pymupdf.open(file["preview_path"]) as pdf:
            assert pdf.page_count == file["page_count"]
        assert sha256(__import__("pathlib").Path(file["preview_path"]).read_bytes()).hexdigest() == file["sha256"]
    assert [sha256(path.read_bytes()).hexdigest() for path in export_task[3]] == before
    bundle.close(created["intent_id"])
    assert not list(bundle.preview_root.glob("*.pdf"))
    with pytest.raises((ExportScopeError, ValueError)):
        bundle.render(created["intent_id"])


def test_duplicate_candidates_keep_scope_rows_and_share_one_rendered_page(duplicate_bundle, duplicate_export_task):
    created = duplicate_bundle.create(request_all(duplicate_export_task, "merged", True))
    assert created["summary"]["selected_count"] == 2
    assert created["summary"]["expected_pages"] == 1
    assert created["total_pages"] == 1
    assert [file["page_count"] for file in created["files"]] == [1]

    register(duplicate_bundle, created)
    rendered = duplicate_bundle.render(created["intent_id"])
    assert rendered["summary"] == created["summary"]
    assert rendered["summary"]["selected_count"] == 2
    assert rendered["summary"]["expected_pages"] == rendered["total_pages"] == 1
    with pymupdf.open(rendered["files"][0]["preview_path"]) as pdf:
        assert pdf.page_count == rendered["total_pages"] == 1

    with duplicate_bundle.journal.locked(created["intent_id"]) as record:
        plan = record.data["plan"]
        assert len(plan["index_rows"]) == 2
        assert len(plan["mappings"]) == 2
        assert {mapping["output_page"] for mapping in plan["mappings"]} == {1}
        assert {mapping["segment_id"] for mapping in plan["mappings"]} == {
            item["id"] for item in duplicate_export_task[4]
        }

    output_root = duplicate_bundle.preview_root.parent / "published-duplicate"
    output_root.mkdir()
    receipt = export_publish.publish_bundle(duplicate_bundle, created["intent_id"], output_root)
    assert receipt["summary"]["selected_count"] == 2
    assert receipt["summary"]["expected_pages"] == receipt["total_pages"] == 1
    assert receipt["row_count"] == 2
    xlsx_path = Path(next(file["path"] for file in receipt["files"] if file["kind"] == "xlsx"))
    import openpyxl

    workbook = openpyxl.load_workbook(xlsx_path, data_only=False)
    try:
        assert workbook["索引"].max_row == 3
        assert workbook["输出映射"].max_row == 3
        assert [workbook["输出映射"][f"F{row}"].value for row in (2, 3)] == [1, 1]
    finally:
        workbook.close()
    duplicate_bundle.close(created["intent_id"])


@pytest.mark.parametrize("mode", ["merged", "by_source", "both"])
def test_real_custom_name_create_render_publish_status_without_preview(bundle, export_task, mode):
    request = request_all(export_task, mode, xlsx=False)
    request["output_name"] = "季度报告.pdf.pdf"
    created = bundle.create(request)
    assert created["output_name"] == "季度报告"
    assert all(file["name"].startswith("季度报告") and file["name"].endswith(".pdf") for file in created["files"])
    register(bundle, created)

    preview = bundle.render(created["intent_id"])
    assert preview["output_name"] == "季度报告"
    assert [file["name"] for file in preview["files"]] == [file["name"] for file in created["files"]]

    output_root = bundle.preview_root.parent / "published-custom"
    output_root.mkdir()
    receipt = export_publish.publish_bundle(bundle, created["intent_id"], output_root)
    assert receipt["state"] == "published"
    assert receipt.get("merged_name") == ("季度报告.pdf" if mode != "by_source" else None)
    final = Path(receipt["directory"])
    assert {path.name for path in final.iterdir()} == {file["name"] for file in receipt["files"]}
    assert (final / "导出清单.json").is_file()

    # Closing removes private previews and compacts the journal, so status must
    # still recognize the custom merged artifact solely from receipt metadata.
    bundle.close(created["intent_id"])
    status = export_publish.status_bundle(bundle, request["job_id"])
    assert status["publication"] == receipt
    assert status["residuals"] == []


def test_render_requires_every_native_preview_registration(bundle, export_task):
    created = bundle.create(request_all(export_task))
    with pytest.raises(ExportScopeError):
        bundle.render(created["intent_id"])
    assert not list(bundle.preview_root.glob("*.pdf"))


def test_changed_scope_cannot_render_and_close_does_not_require_originals(bundle, export_task):
    created = bundle.create(request_all(export_task))
    register(bundle, created)
    record = deepcopy(export_task[4][0])
    record["record_revision"] = 1
    save(export_task, [record])
    with pytest.raises(ExportScopeError):
        bundle.render(created["intent_id"])
    export_task[3][0].unlink()  # generated disposable fixture
    bundle.close(created["intent_id"])


def test_partial_render_keeps_only_owned_files_for_cleanup(bundle, export_task, monkeypatch):
    created = bundle.create(request_all(export_task))
    register(bundle, created)
    original = core._export_pdf_response
    calls = 0
    def fail_second(request):
        nonlocal calls
        calls += 1
        if calls == 2:
            return {"status": "error", "code": "pdf_export_failed", "message": "synthetic failure"}
        return original(request)
    monkeypatch.setattr(core, "_export_pdf_response", fail_second)
    with pytest.raises(ExportScopeError):
        bundle.render(created["intent_id"])
    assert len(list(bundle.preview_root.glob("*.pdf"))) == 1
    bundle.close(created["intent_id"])
    assert not list(bundle.preview_root.glob("*.pdf"))


def test_unknown_and_repeated_tokens_cannot_produce_new_bytes(bundle, export_task):
    with pytest.raises((ExportScopeError, ValueError)):
        bundle.render("../foreign")
    created = bundle.create(request_all(export_task, "merged"))
    register(bundle, created)
    bundle.render(created["intent_id"])
    with pytest.raises(ExportScopeError):
        bundle.render(created["intent_id"])


def test_create_rejects_extra_frontend_paths(bundle, export_task):
    request = request_all(export_task)
    request["selections"] = [{"source_path": "untrusted"}]
    with pytest.raises(ExportScopeError):
        bundle.create(request)


def test_unregistered_previews_are_retained_and_not_marked_closed(bundle, export_task):
    from pathlib import Path
    created = bundle.create(request_all(export_task, "merged"))
    register(bundle, created)
    residual = Path(created["files"][0]["preview_path"])
    residual.write_bytes(b"unregistered foreign sentinel")
    with pytest.raises(ExportScopeError, match="cleanup"):
        bundle.close(created["intent_id"])
    assert residual.read_bytes() == b"unregistered foreign sentinel"
    assert bundle.describe(created["intent_id"])["state"] == "created"


def test_registration_failure_before_first_owner_does_not_block_next_create(bundle, export_task):
    created = bundle.create(request_all(export_task, "merged"))
    bundle.close(created["intent_id"])
    assert bundle.create(request_all(export_task, "merged"))["state"] == "created"


def test_reconcile_closes_only_inactive_private_previews(bundle, export_task):
    created = bundle.create(request_all(export_task, "merged"))
    active_token = created["files"][0]["preview_token"]
    register(bundle, created)
    assert bundle.reconcile([active_token]) == {"residuals": []}
    assert bundle.describe(created["intent_id"])["state"] == "created"
    assert bundle.reconcile([]) == {"residuals": []}
    assert bundle.journal.list_metadata() == []
    assert bundle.create(request_all(export_task, "merged"))["state"] == "created"


@pytest.mark.parametrize("retained", ["preview", "residual_attempts", "previous_attempts", "none"])
def test_receipt_aging_retains_recovery_evidence(bundle, export_task, monkeypatch, retained):
    from engine import export_bundle
    created = bundle.create(request_all(export_task, "merged"))
    bundle.close(created["intent_id"])
    with bundle.journal.locked(created["intent_id"]) as record:
        entry = record.data
        entry["state"] = "published"
        entry["receipt"] = {"intent_id": created["intent_id"]}
        if retained in {"residual_attempts", "previous_attempts"}:
            entry[retained] = [{"path": str(bundle.preview_root / "retained") }]
        record.save(entry)
    if retained == "preview":
        register(bundle, created)
    monkeypatch.setattr(export_bundle, "MAX_RETAINED_RECEIPTS", 0)
    bundle.create(request_all(export_task, "merged"))
    ids = {row["intent_id"] for row in bundle.journal.list_metadata()}
    assert (created["intent_id"] in ids) == (retained != "none")


def test_receipt_aging_retires_on_bytes_before_count_limit(bundle, export_task, monkeypatch):
    monkeypatch.setattr(journal_module, "MAX_JOURNAL_BYTES", 10_000)
    monkeypatch.setattr(journal_module, "MAX_TOTAL_JOURNAL_BYTES", 35_000)
    receipts = [
        _compressed_receipt(bundle, f"2020-01-01T00:00:{index:02d}Z")
        for index in range(20)
    ]
    created = bundle.create(request_all(export_task, "merged"))

    remaining = {row["intent_id"] for row in bundle.journal.list_metadata()
                 if row["job_id"] == "receipt-retention-job"}
    assert len(remaining) < len(receipts)
    assert len(remaining) < 128
    assert str(receipts[-1]["intent_id"]) in remaining
    assert created["state"] == "created"


def test_receipt_aging_keeps_preview_and_residual_receipt_and_output(bundle, export_task, monkeypatch):
    monkeypatch.setattr(journal_module, "MAX_JOURNAL_BYTES", 10_000)
    monkeypatch.setattr(journal_module, "MAX_TOTAL_JOURNAL_BYTES", 35_000)
    blocked = _compressed_receipt(bundle, "2020-01-01T00:00:00Z", blocked=True)
    safe = [
        _compressed_receipt(bundle, f"2020-01-01T00:01:{index:02d}Z")
        for index in range(20)
    ]
    output = Path(blocked["files"][0]["preview_path"])
    before = output.read_bytes()
    created = bundle.create(request_all(export_task, "merged"))

    ids = {row["intent_id"] for row in bundle.journal.list_metadata()}
    assert blocked["intent_id"] in ids
    assert output.read_bytes() == before
    assert created["state"] == "created"
    assert len(ids.intersection({item["intent_id"] for item in safe})) < len(safe)
