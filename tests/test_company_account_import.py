import base64
import json
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from engine.company_account_import import (
    HEADERS, SOURCE_TRUNCATION_NOTICE, AccountImportConflict, AccountImportInvalid,
    _source_text, import_accounts, preview_accounts,
)
from engine.receipt_grouping_store import GroupingStore
from engine.batch_api import handle_batch_request


ACCOUNT = dict(company_name="示例公司", bank_name="示例银行", branch_name="示例支行", account_number="00123456789012345678")


def encoded(rows, header=HEADERS):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "本方账户"
    sheet.append(list(header))
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return base64.b64encode(output.getvalue()).decode()


def test_preview_is_read_only_preserves_digits_and_reports_each_bad_row():
    with GroupingStore() as store:
        result = preview_accounts(store, encoded([
            list(ACCOUNT.values()), list(ACCOUNT.values()),
            ["示例公司", "示例银行", "", 123456789012345678],
            ["=1+1", "示例银行", "", "0001"],
            ["", "示例银行", "", "0002"],
        ]))
        assert result["counts"] == dict(ready=1, duplicate=1, error=3)
        assert result["rows"][0]["account"]["account_number"] == ACCOUNT["account_number"]
        assert result["rows"][0]["source"] == ACCOUNT
        assert result["rows"][1]["source"] == ACCOUNT
        assert result["rows"][3]["source"] == dict(company_name="=1+1", bank_name="示例银行",
                                                   branch_name="", account_number="0001")
        assert result["rows"][4]["source"]["company_name"] == ""
        assert all(row["account"] is None for row in result["rows"][2:])
        assert store.account_list()["total"] == 0


def test_preview_source_preserves_invalid_cells_and_stays_separate_from_normalized_accounts():
    with GroupingStore() as store:
        store.account_save(ACCOUNT)
        conflict = {**ACCOUNT, "company_name": "不同公司"}
        normalized_source = {**ACCOUNT, "company_name": "  示例公司  ", "account_number": "００１２ ３"}
        result = preview_accounts(store, encoded([
            list(normalized_source.values()),
            ["示例公司", "示例银行", None, 12345],
            ["示例公司", "示例银行", None, "=1+1"],
            list(conflict.values()),
        ]))
        assert result["counts"] == dict(ready=1, duplicate=0, error=3)
        assert result["rows"][0]["source"] == normalized_source
        assert result["rows"][0]["account"]["company_name"] == ACCOUNT["company_name"]
        assert result["rows"][0]["account"]["account_number"] == "00123"
        assert result["rows"][1]["source"]["account_number"] == "12345"
        assert result["rows"][1]["source"]["branch_name"] == ""
        assert result["rows"][1]["account"] is None
        assert result["rows"][2]["source"]["account_number"] == "=1+1"
        assert result["rows"][2]["account"] is None
        assert result["rows"][3]["account"] == conflict
        assert result["rows"][3]["source"] == conflict
        assert store.account_list()["total"] == 1


@pytest.mark.parametrize(("value", "expected"), [
    (None, ""), (0, "0"), (False, "False"), ("  原文\n", "  原文\n"), ("原\x00文", "原�文"),
])
def test_source_text_uses_visible_safe_text(value, expected):
    assert _source_text(value) == expected


def test_preview_marks_long_source_cells_without_changing_account_validation():
    with GroupingStore() as store:
        result = preview_accounts(store, encoded([["🧾" * 257] * 4]))
        row = result["rows"][0]
        assert row["status"] == "error"
        assert row["account"] is None
        assert all(value == "🧾" * 256 + SOURCE_TRUNCATION_NOTICE for value in row["source"].values())
        assert store.account_list()["total"] == 0


def test_maximum_valid_preview_with_source_stays_within_response_budget():
    widest_account = {key: "🧾" * 256 for key in ("company_name", "bank_name", "branch_name")}
    widest_account["account_number"] = "0" * 128
    with GroupingStore() as store:
        result = preview_accounts(store, encoded([list(widest_account.values())] * 500))
        assert len(result["rows"]) == 500
        assert result["rows"][-1]["row_number"] == 501
        assert result["counts"] == dict(ready=1, duplicate=499, error=0)
        response_bytes = json.dumps({"status": "ok", "data": result}, ensure_ascii=False).encode("utf-8")
        assert len(response_bytes) < 4 * 1024 * 1024


def test_commit_rechecks_conflicts_rolls_back_and_retry_is_idempotent():
    with GroupingStore() as store:
        store.account_save(ACCOUNT, active=False)
        assert import_accounts(store, [ACCOUNT]) == dict(created=0, skipped=1)
        new = {**ACCOUNT, "account_number": "0099"}
        with pytest.raises(AccountImportConflict):
            import_accounts(store, [new, {**ACCOUNT, "company_name": "不同公司"}])
        assert store.account_list(active_only=False)["total"] == 1
        assert import_accounts(store, [new, new]) == dict(created=1, skipped=1)
        assert import_accounts(store, [new]) == dict(created=0, skipped=1)
        assert next(item for item in store.account_list(active_only=False)["items"]
                    if item["account_number"] == ACCOUNT["account_number"])["active"] is False


@pytest.mark.parametrize("payload", ["not a workbook", "A" * 700000, base64.b64encode(b"invalid ZIP").decode()], ids=["not_base64", "oversized", "not_zip"])
def test_bad_workbook_is_rejected(payload):
    with GroupingStore() as store, pytest.raises(AccountImportInvalid):
        preview_accounts(store, payload)


def test_header_row_limit_extra_columns_and_external_content_rejected():
    with GroupingStore() as store:
        for payload in (encoded([], header=("错误列", *HEADERS[1:])),
                        encoded([list(ACCOUNT.values())] * 501),
                        encoded([[*ACCOUNT.values(), "额外列"]])):
            with pytest.raises(AccountImportInvalid):
                preview_accounts(store, payload)
        raw = BytesIO(base64.b64decode(encoded([list(ACCOUNT.values())])))
        with ZipFile(raw, "a") as archive:
            archive.writestr("xl/externalLinks/externalLink1.xml", "<externalLink/>")
        with pytest.raises(AccountImportInvalid):
            preview_accounts(store, base64.b64encode(raw.getvalue()).decode())


@pytest.mark.parametrize("encoding", ["utf-8", "utf-16", "utf-16-le", "utf-32"])
def test_xml_entities_are_rejected_before_parsing_in_any_encoding(encoding):
    raw = BytesIO(base64.b64decode(encoded([list(ACCOUNT.values())])))
    output = BytesIO()
    with ZipFile(raw) as source, ZipFile(output, "w") as target:
        for entry in source.infolist():
            content = source.read(entry)
            if entry.filename == "xl/worksheets/sheet1.xml":
                text = '<!DOCTYPE worksheet [<!ENTITY demo "unexpected">]>' + content.decode("utf-8")
                content = text.encode(encoding)
            target.writestr(entry, content)
    with GroupingStore() as store, pytest.raises(AccountImportInvalid):
        preview_accounts(store, base64.b64encode(output.getvalue()).decode())


def test_batch_dispatch_and_template_contract(tmp_path):
    paths = dict(database_path=str(tmp_path / "batch.sqlite3"), grouping_database_path=str(tmp_path / "accounts.sqlite3"))
    preview = handle_batch_request({"op": "batch_counterparty_account_import_preview", **paths,
                                    "workbook_base64": encoded([list(ACCOUNT.values())])})
    assert preview["status"] == "ok"
    commit = handle_batch_request({"op": "batch_counterparty_account_import", **paths, "accounts": [ACCOUNT]})
    assert commit == {"status": "ok", "data": {"created": 1, "skipped": 0}}
    template = Path(__file__).resolve().parents[1] / "public/templates/本方账户导入模板.xlsx"
    book = load_workbook(template)
    sheet = book["本方账户"]
    assert tuple(cell.value for cell in sheet[1]) == HEADERS
    assert all(sheet.cell(row, 4).number_format == "@" for row in (2, 250, 501))
    assert all(sheet.cell(2, col).value is None for col in range(1, 5))
    book.close()
